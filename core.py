"""Núcleo do conversor Excel -> Firebird (sem interface gráfica)."""
import os
import sys
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

def app_dir():
    """Pasta do .exe (ou do script, quando roda com python)."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


# ---------------------------------------------------------------- Firebird
# Driver firebirdsql: Python puro, fala direto com o servidor pela rede.
# Não precisa de fbclient.dll nem de nada instalado. Funciona com 2.5, 3, 4 e 5.
def open_conn(p):
    """p: dict com host, port, db, user, password, charset."""
    import firebirdsql
    import passlib.hash  # noqa: F401  (login antigo do Firebird 2.5 / Legacy_Auth)

    host = (p.get("host") or "").strip() or "localhost"
    port = int((p.get("port") or "").strip() or 3050)
    db = p["db"].strip()
    base = dict(host=host, port=port, database=db, user=p["user"].strip(),
                password=p["password"], charset=p.get("charset") or "WIN1252")
    # 1) padrão SRP + criptografia (Firebird 3+)  2) SRP sem criptografia
    # 3) login antigo (Firebird 2.5 ou servidor configurado com Legacy_Auth)
    tentativas = [{}, {"wire_crypt": False}, {"auth_plugin_name": "Legacy_Auth", "wire_crypt": False}]
    erros = []
    for extra in tentativas:
        try:
            return firebirdsql.connect(**base, **extra)
        except Exception as e:  # noqa: BLE001
            erros.append(e)
            msg = str(e).lower()
            # erro de rede: não adianta tentar outro modo de login
            if any(t in msg for t in ("refused", "timed out", "getaddrinfo", "unreachable",
                                      "no route", "11001", "10061", "10060")):
                break
    raise ConnectionError(_explica(erros, host, port, db))


def _explica(erros, host, port, db):
    textos = [str(e) for e in erros]
    low = " | ".join(textos).lower()
    m = " | ".join(dict.fromkeys(textos))
    if "refused" in low or "10061" in low:
        dica = f"Servidor Firebird não respondeu em {host}:{port}. Confira o host, a porta e se o serviço está rodando."
    elif "timed out" in low or "10060" in low or "getaddrinfo" in low or "11001" in low:
        dica = f"Não achei o servidor {host}. Confira o host/IP e a rede."
    elif "i/o error" in low or "no such file" in low or "335544344" in low:
        dica = f"Banco não encontrado no servidor: {db}. O caminho é o do .fdb na máquina do servidor."
    elif any(t in low for t in ("password", "user name", "login", "unauthorized", "335544472")):
        dica = "Usuário ou senha incorretos."
    else:
        dica = "Não foi possível conectar."
    return f"{dica}\n\nDetalhe: {m}"


def list_tables(conn):
    cur = conn.cursor()
    cur.execute(
        "SELECT TRIM(RDB$RELATION_NAME) FROM RDB$RELATIONS "
        "WHERE COALESCE(RDB$SYSTEM_FLAG, 0) = 0 AND RDB$VIEW_BLR IS NULL ORDER BY 1"
    )
    return [r[0] for r in cur.fetchall()]


def kind_of(ftype, subtype, scale):
    """Converte o tipo do Firebird num 'kind' simples."""
    scale = scale or 0
    if ftype in (7, 8, 16):
        return f"dec:{-scale}" if scale < 0 else "int"
    if ftype in (10, 27):
        return "float"
    if ftype in (14, 37):
        return "str"
    if ftype == 12:
        return "date"
    if ftype == 35:
        return "ts"
    return "raw"


def list_columns(conn, table):
    """Devolve dict ordenado {coluna: kind}."""
    cur = conn.cursor()
    cur.execute(
        "SELECT TRIM(rf.RDB$FIELD_NAME), f.RDB$FIELD_TYPE, "
        "COALESCE(f.RDB$FIELD_SUB_TYPE, 0), COALESCE(f.RDB$FIELD_SCALE, 0) "
        "FROM RDB$RELATION_FIELDS rf "
        "JOIN RDB$FIELDS f ON f.RDB$FIELD_NAME = rf.RDB$FIELD_SOURCE "
        "WHERE rf.RDB$RELATION_NAME = ? ORDER BY rf.RDB$FIELD_POSITION",
        (table,),
    )
    return {r[0]: kind_of(r[1], r[2], r[3]) for r in cur.fetchall()}


# ------------------------------------------------------------------- Excel
def excel_sheets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_headers(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        first = next(wb[sheet].iter_rows(min_row=1, max_row=1, values_only=True), ())
    finally:
        wb.close()
    return [str(h).strip() if h not in (None, "") else f"Coluna{i + 1}" for i, h in enumerate(first)]


def _zero_fmt_width(fmt):
    """'00000000' -> 8. Formato que não é só zeros -> 0."""
    if not fmt:
        return 0
    fmt = fmt.split(";")[0].replace('"', "").strip()
    return len(fmt) if fmt and set(fmt) == {"0"} else 0


def _cell_value(cell):
    """Valor da célula. Número com formato de zeros à esquerda (ex: 00000000)
    vira texto já com os zeros, pra não perder o que o Excel mostra na tela."""
    v = cell.value
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return v
    w = _zero_fmt_width(getattr(cell, "number_format", None))
    if w and float(v).is_integer() and v >= 0:
        return str(int(v)).zfill(w)
    return v


def read_excel(path, sheet):
    """Devolve (headers, [(linha_excel, tupla_valores), ...]). Cabeçalho na linha 1."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb[sheet].iter_rows()
        first = [c.value for c in next(it, ())]
        headers = [str(h).strip() if h not in (None, "") else f"Coluna{i + 1}" for i, h in enumerate(first)]
        n = len(headers)
        rows = []
        for line, cells in enumerate(it, start=2):
            row = tuple(_cell_value(c) for c in cells[:n])
            row = row + (None,) * max(0, n - len(row))
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue
            rows.append((line, row))
    finally:
        wb.close()
    return headers, rows


# --------------------------------------------------------------- Conversão
def _num(s):
    s = s.strip().replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    return s


def _parse_date(s, formatos):
    for f in formatos:
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    raise ValueError(f"data inválida: {s!r}")


def convert(v, kind):
    """Converte o valor do Excel pro tipo da coluna. Vazio vira None."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v == "":
            return None

    if kind == "int":
        n = Decimal(_num(v)) if isinstance(v, str) else Decimal(str(v))
        if n != n.to_integral_value():
            raise ValueError(f"valor não inteiro: {v!r}")
        return int(n)
    if kind.startswith("dec:"):
        casas = int(kind.split(":")[1])
        n = Decimal(_num(v)) if isinstance(v, str) else Decimal(str(v))
        return n.quantize(Decimal(1).scaleb(-casas), rounding=ROUND_HALF_UP)
    if kind == "float":
        return float(_num(v)) if isinstance(v, str) else float(v)
    if kind == "str":
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    if kind == "date":
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        return _parse_date(v, ("%d/%m/%Y", "%Y-%m-%d")).date()
    if kind == "ts":
        if isinstance(v, datetime):
            return v
        if isinstance(v, date):
            return datetime.combine(v, datetime.min.time())
        return _parse_date(v, ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
                               "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"))
    return v


# -------------------------------------------------------------- Checagens
def _perde_zero(v):
    """Texto só de dígitos com zero na frente, ex: '0300600'."""
    return isinstance(v, str) and len(v.strip()) > 1 and v.strip().isdigit() and v.strip().startswith("0")


def check_risks(headers, rows, pairs, kinds, pads=None):
    """Procura situações em que o zero à esquerda se perde.
    Devolve (bloqueios, avisos): listas de textos."""
    pads = pads or {}
    bloqueios, avisos = [], []
    for x, d in pairs:
        i = headers.index(x)
        kind = kinds[d]
        com_zero = [(ln, r[i]) for ln, r in rows if _perde_zero(r[i])]
        if kind != "str" and com_zero:
            ln, v = com_zero[0]
            bloqueios.append(
                f'A coluna "{d}" do banco é numérica ({kind.split(":")[0]}) e não guarda zero à esquerda: '
                f'{len(com_zero)} valor(es) da coluna "{x}" do Excel têm zero na frente '
                f'(ex: linha {ln}: "{v.strip()}" ficaria {int(v.strip())}). '
                f'Mude a coluna "{d}" para VARCHAR no banco, ou tire essa coluna da atualização.')
        if kind == "str" and not pads.get(d):
            soltos = [(ln, r[i]) for ln, r in rows
                      if isinstance(r[i], (int, float)) and not isinstance(r[i], bool)]
            if soltos:
                ln, v = soltos[0]
                avisos.append(
                    f'Atenção: "{x}" vem como NÚMERO no Excel em {len(soltos)} linha(s) (ex: linha {ln}: {v:g}). '
                    f'Se esses códigos tinham zero na frente, preencha "Zeros" ao lado de "{d}" '
                    f'(ex: 7 para CEST) ou formate a coluna do Excel como Texto.')
    return bloqueios, avisos


# -------------------------------------------------------------- Atualização
def run_update(conn, table, ref_x, ref_db, pairs, headers, rows, kinds,
               write_null=False, dry=True, log=print, progress=lambda a, b: None,
               max_log=300, pads=None):
    """
    pairs: [(coluna_excel, coluna_banco), ...]
    pads: {coluna_banco: tamanho} completa com zeros à esquerda (só colunas de texto).
    Faz UPDATE "table" SET col = ? WHERE ref_db = ? para cada linha do Excel.
    dry=True: executa e dá rollback no fim (só conta o que seria alterado).
    """
    pads = pads or {}

    def conv(v, col):
        r = convert(v, kinds[col])
        w = pads.get(col)
        if w and kinds[col] == "str" and isinstance(r, str) and r.isdigit():
            r = r.zfill(w)
        return r

    ix_ref = headers.index(ref_x)
    ix_vals = [headers.index(x) for x, _ in pairs]
    sets = ", ".join(f'"{d}" = ?' for _, d in pairs)
    sql = f'UPDATE "{table}" SET {sets} WHERE "{ref_db}" = ?'

    st = {"total": len(rows), "atualizadas": 0, "nao_encontradas": 0,
          "puladas": 0, "erros": 0, "duplicadas": 0, "divergentes": 0}
    cur = conn.cursor()
    mostrados = [0]

    def aviso(msg):
        if mostrados[0] < max_log:
            log(msg)
        elif mostrados[0] == max_log:
            log("... (demais avisos omitidos, veja o resumo no final)")
        mostrados[0] += 1

    # amostra do que será gravado (ajuda a achar de onde some o zero à esquerda)
    rotulo = {"int": "inteiro", "float": "decimal", "str": "texto", "date": "data", "ts": "data/hora"}
    log("Tipo das colunas do banco: " + ", ".join(
        f"{d}={'decimal' if kinds[d].startswith('dec') else rotulo.get(kinds[d], kinds[d])}" for _, d in pairs))
    mostradas = 0
    for line, row in rows:
        if mostradas >= 3:
            break
        try:
            if any(row[i] is not None for i in ix_vals):
                for i, (x, d) in zip(ix_vals, pairs):
                    orig = row[i]
                    log(f"Amostra linha {line}: {x} lido do Excel = {orig!r} ({type(orig).__name__})"
                        f"  ->  vai gravar em {d} = {conv(orig, d)!r}")
                mostradas += 1
        except Exception:
            pass

    try:
        for n, (line, row) in enumerate(rows, 1):
            try:
                key = conv(row[ix_ref], ref_db)
                if key is None:
                    st["puladas"] += 1
                    aviso(f"Linha {line}: referência vazia, pulada")
                    continue
                vals = [conv(row[i], d) for i, (_, d) in zip(ix_vals, pairs)]
                if not write_null and any(v is None for v in vals):
                    st["puladas"] += 1
                    aviso(f"Linha {line} (ref {key}): célula vazia, pulada")
                    continue
                cur.execute(sql, vals + [key])
                cnt = cur.rowcount
                if cnt == 0:
                    st["nao_encontradas"] += 1
                    aviso(f"Linha {line}: ref {key} não encontrada no banco")
                else:
                    st["atualizadas"] += 1
                    # conferência: lê de volta os códigos com zero na frente e compara
                    for v, (_, d) in zip(vals, pairs):
                        if _perde_zero(v):
                            cur.execute(f'SELECT "{d}" FROM "{table}" WHERE "{ref_db}" = ?', (key,))
                            got = cur.fetchone()
                            got = got[0] if got else None
                            if got is None or str(got).rstrip() != v:
                                st["divergentes"] += 1
                                aviso(f'Linha {line} (ref {key}): gravei "{v}" em {d}, mas o banco ficou com "{got}"')
                    if cnt and cnt > 1:
                        st["duplicadas"] += 1
                        aviso(f"Linha {line}: ref {key} atualizou {cnt} registros (referência não é única)")
            except Exception as e:  # erro de conversão ou do banco: segue pra próxima
                st["erros"] += 1
                aviso(f"Linha {line}: erro - {e}")
            if n % 50 == 0 or n == len(rows):
                progress(n, len(rows))
        if dry:
            conn.rollback()
        else:
            conn.commit()
    except Exception:
        conn.rollback()
        raise
    return st
